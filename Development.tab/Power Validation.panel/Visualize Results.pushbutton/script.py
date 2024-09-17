#! python3

import clr
import os
import sys
import traceback
import math
import openpyxl
from System.Collections.Generic import List

# Add Revit references
clr.AddReference('RevitAPI')
clr.AddReference('RevitAPIUI')

from Autodesk.Revit.DB import *
from Autodesk.Revit.UI import *
from Autodesk.Revit.UI.Selection import ObjectType
from System import Guid

# Get the Revit application and document
uidoc = __revit__.ActiveUIDocument
doc = uidoc.Document

def get_element_transforms(element, host_doc):
    """
    Get the cumulative transformations of an element to the host document's coordinate system.
    Returns a list of (transform, link_instance_chain) tuples.
    """
    if element.Document.Equals(host_doc):
        # Element is in the host document
        return [(Transform.Identity, [])]
    else:
        # Element is in a linked document
        transforms = []
        link_instances = FilteredElementCollector(host_doc).OfClass(RevitLinkInstance).ToElements()
        for link_instance in link_instances:
            linked_doc = link_instance.GetLinkDocument()
            if linked_doc:
                # If the linked document matches the element's document
                if linked_doc.Equals(element.Document):
                    # Found a direct link
                    transform = link_instance.GetTransform()
                    transforms.append((transform, [link_instance]))
                else:
                    # Check for nested links
                    nested_transforms = get_element_transforms_in_link(
                        element, linked_doc, link_instance.GetTransform(), [link_instance])
                    transforms.extend(nested_transforms)
        return transforms

def get_element_transforms_in_link(element, current_doc, current_transform, link_instance_chain):
    """
    Recursive helper function to get the cumulative transformations of an element in nested links.
    Returns a list of (transform, link_instance_chain) tuples.
    """
    if element.Document.Equals(current_doc):
        # Element is in the current document
        return [(current_transform, link_instance_chain)]
    else:
        # Search nested links in current_doc
        transforms = []
        link_instances = FilteredElementCollector(current_doc).OfClass(RevitLinkInstance).ToElements()
        for link_instance in link_instances:
            linked_doc = link_instance.GetLinkDocument()
            if linked_doc:
                nested_transform = current_transform.Multiply(link_instance.GetTransform())
                nested_chain = link_instance_chain + [link_instance]
                if linked_doc.Equals(element.Document):
                    # Found the linked document
                    transforms.append((nested_transform, nested_chain))
                else:
                    # Recursively search deeper
                    nested_transforms = get_element_transforms_in_link(
                        element, linked_doc, nested_transform, nested_chain)
                    transforms.extend(nested_transforms)
        return transforms

def TransformBoundingBox(bbox, transform):
    """
    Transforms a bounding box using the given transform.
    """
    corners = [
        XYZ(bbox.Min.X, bbox.Min.Y, bbox.Min.Z),
        XYZ(bbox.Min.X, bbox.Min.Y, bbox.Max.Z),
        XYZ(bbox.Min.X, bbox.Max.Y, bbox.Min.Z),
        XYZ(bbox.Min.X, bbox.Max.Y, bbox.Max.Z),
        XYZ(bbox.Max.X, bbox.Min.Y, bbox.Min.Z),
        XYZ(bbox.Max.X, bbox.Min.Y, bbox.Max.Z),
        XYZ(bbox.Max.X, bbox.Max.Y, bbox.Min.Z),
        XYZ(bbox.Max.X, bbox.Max.Y, bbox.Max.Z),
    ]
    transformed_corners = [transform.OfPoint(corner) for corner in corners]
    xs = [pt.X for pt in transformed_corners]
    ys = [pt.Y for pt in transformed_corners]
    zs = [pt.Z for pt in transformed_corners]
    transformed_bbox = BoundingBoxXYZ()
    transformed_bbox.Min = XYZ(min(xs), min(ys), min(zs))
    transformed_bbox.Max = XYZ(max(xs), max(ys), max(zs))
    return transformed_bbox

def UnionBoundingBoxes(bbox1, bbox2):
    """
    Returns the union of two bounding boxes.
    """
    min_point = XYZ(
        min(bbox1.Min.X, bbox2.Min.X),
        min(bbox1.Min.Y, bbox2.Min.Y),
        min(bbox1.Min.Z, bbox2.Min.Z)
    )
    max_point = XYZ(
        max(bbox1.Max.X, bbox2.Max.X),
        max(bbox1.Max.Y, bbox2.Max.Y),
        max(bbox1.Max.Z, bbox2.Max.Z)
    )
    union_bbox = BoundingBoxXYZ()
    union_bbox.Min = min_point
    union_bbox.Max = max_point
    return union_bbox

def adjust_view(view, combined_bbox, padding=5):
    """
    Adjust the view's section box and orientation to focus on the given combined bounding box.
    """
    # Apply padding
    min_point = XYZ(
        combined_bbox.Min.X - padding,
        combined_bbox.Min.Y - padding,
        combined_bbox.Min.Z - padding
    )
    max_point = XYZ(
        combined_bbox.Max.X + padding,
        combined_bbox.Max.Y + padding,
        combined_bbox.Max.Z + padding
    )

    # Set the section box
    section_box = BoundingBoxXYZ()
    section_box.Min = min_point
    section_box.Max = max_point
    view.IsSectionBoxActive = True
    view.SectionBox = section_box

    # Adjust the view orientation
    center = min_point.Add(max_point).Multiply(0.5)
    view_direction = XYZ(0, 1, -0.5).Normalize()  # Adjusted for better viewing angle
    up_direction = XYZ.BasisZ  # Z-axis as up

    # Position the camera
    bbox_diagonal = max_point.DistanceTo(min_point)
    distance = bbox_diagonal * 1.5  # Adjust multiplier as needed
    eye_position = center.Subtract(view_direction.Multiply(distance))

    # Create and set the new orientation
    forward_direction = center.Subtract(eye_position).Normalize()

    # Ensure that up_direction and forward_direction are not parallel
    dot = abs(forward_direction.DotProduct(up_direction))
    if dot >= 0.99:
        # If they are nearly parallel, choose a different upDirection
        up_direction = XYZ.BasisX

    right_direction = up_direction.CrossProduct(forward_direction).Normalize()
    up_direction = forward_direction.CrossProduct(right_direction).Normalize()

    # Verify up_direction is not zero-length after recalculating
    if up_direction.IsZeroLength():
        print("upDirection has zero length after adjustment. Setting to default XYZ.BasisZ.")
        up_direction = XYZ.BasisZ

    orientation = ViewOrientation3D(eye_position, up_direction, forward_direction)
    view.SetOrientation(orientation)

def main():
    # Prompt the user to select the Excel file
    from System.Windows.Forms import OpenFileDialog, DialogResult
    open_dialog = OpenFileDialog()
    open_dialog.Filter = "Excel Files (*.xlsx)|*.xlsx"
    open_dialog.Title = "Select the Excel Report File"
    if open_dialog.ShowDialog() != DialogResult.OK:
        print("No file selected. Exiting.")
        return
    excel_file_path = open_dialog.FileName

    # Open the Excel file
    try:
        wb = openpyxl.load_workbook(excel_file_path)
        ws = wb.active
    except Exception as e:
        TaskDialog.Show("Error", f"Failed to open Excel file: {e}")
        return

    # Prepare data from Excel
    data_rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        data = {
            'source_unique_id': row[0],
            'target_unique_id': row[7],
            'distance': row[14],
        }
        data_rows.append(data)

    if not data_rows:
        TaskDialog.Show("Information", "No data found in the Excel file.")
        return

    # Get or create a 3D view
    view_name = "Temporary Visualization View"
    view = None
    for v in FilteredElementCollector(doc).OfClass(View3D):
        if not v.IsTemplate and v.Name == view_name:
            view = v
            break

    if not view:
        view_family_type = next((vft for vft in FilteredElementCollector(doc).OfClass(ViewFamilyType)
                                 if vft.ViewFamily == ViewFamily.ThreeDimensional), None)
        if not view_family_type:
            print("No 3D view family type found.")
            return

        t = Transaction(doc, "Create Temporary View")
        t.Start()
        try:
            view = View3D.CreateIsometric(doc, view_family_type.Id)
            view.Name = view_name
            t.Commit()
        except Exception as e:
            t.RollBack()
            print(f"Failed to create view: {e}")
            return

    index = 0
    while index < len(data_rows):
        data = data_rows[index]
        source_unique_id = data['source_unique_id']
        target_unique_id = data['target_unique_id']
        distance = data['distance']

        # Find the source element
        source_element = doc.GetElement(source_unique_id)
        if not source_element:
            # Search in linked documents
            for link_instance in FilteredElementCollector(doc).OfClass(RevitLinkInstance):
                linked_doc = link_instance.GetLinkDocument()
                if linked_doc:
                    source_element = linked_doc.GetElement(source_unique_id)
                    if source_element:
                        break
        if not source_element:
            print(f"Source element with UniqueId {source_unique_id} not found.")
            index += 1
            continue

        # Find the target element
        target_element = doc.GetElement(target_unique_id)
        if not target_element:
            # Search in linked documents
            for link_instance in FilteredElementCollector(doc).OfClass(RevitLinkInstance):
                linked_doc = link_instance.GetLinkDocument()
                if linked_doc:
                    target_element = linked_doc.GetElement(target_unique_id)
                    if target_element:
                        break
        if not target_element:
            print(f"Target element with UniqueId {target_unique_id} not found.")
            index += 1
            continue

        # Get transformations for source element
        source_transforms = get_element_transforms(source_element, doc)
        if not source_transforms:
            print(f"Could not find link instance for source element {source_element.Id}")
            index += 1
            continue

        # Get transformations for target element
        target_transforms = get_element_transforms(target_element, doc)
        if not target_transforms:
            print(f"Could not find link instance for target element {target_element.Id}")
            index += 1
            continue

        # Collect all transformed bounding boxes
        transformed_bounding_boxes = []
        for source_transform, _ in source_transforms:
            source_bbox = source_element.get_BoundingBox(None)
            if source_bbox:
                transformed_source_bbox = TransformBoundingBox(source_bbox, source_transform)
            else:
                print(f"No bounding box for source element {source_element.Id}")
                continue

            for target_transform, _ in target_transforms:
                target_bbox = target_element.get_BoundingBox(None)
                if target_bbox:
                    transformed_target_bbox = TransformBoundingBox(target_bbox, target_transform)
                else:
                    print(f"No bounding box for target element {target_element.Id}")
                    continue

                # Combine the two transformed bounding boxes
                combined_bbox = UnionBoundingBoxes(transformed_source_bbox, transformed_target_bbox)
                transformed_bounding_boxes.append(combined_bbox)

        # Get the union of all combined bounding boxes
        if transformed_bounding_boxes:
            total_combined_bbox = transformed_bounding_boxes[0]
            for bbox in transformed_bounding_boxes[1:]:
                total_combined_bbox = UnionBoundingBoxes(total_combined_bbox, bbox)
        else:
            print("No bounding boxes found for the elements.")
            index += 1
            continue

        # Adjust the view
        t = Transaction(doc, "Adjust View")
        t.Start()
        try:
            adjust_view(view, total_combined_bbox)
            t.Commit()
        except Exception as e:
            t.RollBack()
            print(f"Failed to adjust view: {e}")
            index += 1
            continue

        # Activate the view
        uidoc.ActiveView = view

        # Display information to the user
        dialog = TaskDialog("Element Pair Visualization")
        dialog.MainInstruction = f"Pair {index + 1} of {len(data_rows)}"
        dialog.MainContent = f"Distance: {distance} m"
        if index < len(data_rows) - 1:
            dialog.AddCommandLink(TaskDialogCommandLinkId.CommandLink1, "Next")
        if index > 0:
            dialog.AddCommandLink(TaskDialogCommandLinkId.CommandLink2, "Previous")
        dialog.CommonButtons = TaskDialogCommonButtons.Close
        dialog.DefaultButton = TaskDialogResult.Close

        result = dialog.Show()

        if result == TaskDialogResult.CommandLink1:
            index += 1
        elif result == TaskDialogResult.CommandLink2:
            index -= 1
        else:
            break

    print("Visualization completed.")

if __name__ == '__main__':
    main()
